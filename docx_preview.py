#!/usr/bin/env python3
"""
Script to preview the content of a DOCX document
"""

from docx import Document
import os


def preview_docx(file_path):
    """
    Preview the content of a DOCX file
    """
    try:
        # Load the DOCX document
        doc = Document(file_path)
        
        print(f"Previewing document: {file_path}")
        print("="*50)
        
        # Extract and print paragraphs
        for i, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():  # Only print non-empty paragraphs
                print(f"{i+1}. {paragraph.text}")
        
        print("="*50)
        
        # Also check for tables
        if doc.tables:
            print("\nTables found in the document:")
            for i, table in enumerate(doc.tables):
                print(f"\nTable {i+1}:")
                for j, row in enumerate(table.rows):
                    cells_text = [cell.text for cell in row.cells]
                    print(f"  Row {j+1}: {cells_text}")
        
        # Check for other elements
        print(f"\nDocument statistics:")
        print(f"- Number of paragraphs: {len(doc.paragraphs)}")
        print(f"- Number of tables: {len(doc.tables)}")
        
    except Exception as e:
        print(f"Error reading the document: {e}")


def main():
    # Path to the DOCX file in the workspace
    docx_file = "/workspace/для запуска.docx"
    
    # Check if the file exists
    if not os.path.exists(docx_file):
        print(f"File {docx_file} does not exist!")
        return
    
    # Preview the DOCX file
    preview_docx(docx_file)


if __name__ == "__main__":
    main()